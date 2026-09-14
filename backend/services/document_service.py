from io import BytesIO

from docx import Document
from pypdf import PdfReader
from werkzeug.datastructures import FileStorage

from services.base_service import BaseFileProcessor
from services.translator import TextTranslator
from utils.exceptions import AppError, ValidationError


class DocumentService(BaseFileProcessor):
    """Extracts and translates PDF, DOCX and TXT documents."""

    category = "document"
    MAX_EXTRACTED_CHARACTERS = 30_000
    CHUNK_SIZE = 7_000

    def __init__(self, translator: TextTranslator, file_validator=None):
        super().__init__(file_validator)
        self._translator = translator

    def process(
        self,
        uploaded_file: FileStorage,
        source_language: str,
        target_language: str,
    ) -> dict:
        metadata = self.validate_file(uploaded_file)
        data = self.read_bytes(uploaded_file)

        try:
            extractors = {
                "pdf": self._extract_pdf,
                "docx": self._extract_docx,
                "txt": self._extract_txt,
            }
            original_text = extractors[metadata.extension](data).strip()
        except ValidationError:
            raise
        except Exception as exc:
            raise AppError(
                "DOCUMENT_PROCESSING_ERROR",
                "No fue posible leer el documento.",
                422,
            ) from exc

        if not original_text:
            raise ValidationError(
                "NO_DOCUMENT_TEXT",
                "El documento no contiene texto procesable.",
            )
        if len(original_text) > self.MAX_EXTRACTED_CHARACTERS:
            raise ValidationError(
                "DOCUMENT_TOO_LONG",
                "El documento supera el límite de 30,000 caracteres.",
            )

        translated_chunks = [
            self._translator.translate(
                chunk, source_language, target_language
            )
            for chunk in self._split_text(original_text)
        ]
        return {
            "filename": metadata.filename,
            "original": original_text,
            "translation": "\n\n".join(translated_chunks),
        }

    @staticmethod
    def _extract_pdf(data: bytes) -> str:
        if not data.startswith(b"%PDF"):
            raise ValidationError(
                "INVALID_FORMAT", "El archivo no parece ser un PDF válido."
            )
        reader = PdfReader(BytesIO(data))
        pages = []
        for number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                pages.append(f"[Página {number}]\n{text}")
        return "\n\n".join(pages)

    @staticmethod
    def _extract_docx(data: bytes) -> str:
        document = Document(BytesIO(data))
        sections = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        for table_number, table in enumerate(document.tables, start=1):
            rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                sections.append(f"[Tabla {table_number}]\n" + "\n".join(rows))
        return "\n\n".join(sections)

    @staticmethod
    def _extract_txt(data: bytes) -> str:
        try:
            return data.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                return data.decode("latin-1")
            except UnicodeDecodeError as exc:
                raise ValidationError(
                    "INVALID_ENCODING", "El archivo TXT debe utilizar UTF-8."
                ) from exc

    def _split_text(self, text: str) -> list[str]:
        chunks: list[str] = []
        current = ""
        for paragraph in text.split("\n\n"):
            paragraph = paragraph.strip()
            if not paragraph:
                continue
            if len(paragraph) > self.CHUNK_SIZE:
                if current:
                    chunks.append(current)
                    current = ""
                chunks.extend(
                    paragraph[start : start + self.CHUNK_SIZE]
                    for start in range(0, len(paragraph), self.CHUNK_SIZE)
                )
                continue
            candidate = f"{current}\n\n{paragraph}" if current else paragraph
            if len(candidate) <= self.CHUNK_SIZE:
                current = candidate
            else:
                chunks.append(current)
                current = paragraph
        if current:
            chunks.append(current)
        return chunks

